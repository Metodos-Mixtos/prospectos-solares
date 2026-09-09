"""Exporta el PDF a un PPTX con los elementos sueltos y editables.

A diferencia de powerpoint.py, que rasteriza cada lamina entera y garantiza
fidelidad exacta, este exporta cada pieza por separado:

  - el texto de la lamina, como cuadros de texto reales, editables y movibles
  - cada figura, como una imagen suelta que se puede mover y redimensionar
  - los filetes y las lineas, como formas

Como se distingue una cosa de otra. El texto que compone LaTeX sale con la
tipografia NimbusSanL; el que dibuja matplotlib dentro de una figura sale con
Arial. Esa diferencia separa sin ambiguedad el texto de la lamina del interior
de una figura, y evita convertir los rotulos de un grafico en cien cuadros de
texto sueltos.

Lo que se pierde, y hay que decirlo: PowerPoint no compone como LaTeX. Los
cortes de linea y el interlineado quedan parecidos, no identicos. El interior de
los graficos sigue siendo una imagen.

Uso:  python powerpoint_editable.py [pdf] [pptx] [n1 n2 ...]
"""
import datetime
import pathlib
import sys

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import pptx_comun

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PDF = RAIZ / "outputs" / "presentacion_sb" / "build" / "prospectos_solares_sb.pdf"
SAL = RAIZ / "entregables" / "presentacion_prospectos_solares_editable.pptx"
TMP = RAIZ / "outputs" / "presentacion_sb" / "_pptx_editable"

#: 16:9 panoramico, el mismo tamano que la version rasterizada.
PULGADAS = (13.333, 7.5)
#: Resolucion a la que se rasteriza cada figura suelta.
PPP_FIGURA = 400
#: LaTeX compone en Helvetica; Arial es su equivalente metrico y esta en
#: cualquier PowerPoint.
TIPO = "Arial"
#: Holgura de la caja, en puntos del PDF. Arial mide algo distinto de la
#: Helvetica que compone LaTeX, asi que sin holgura PowerPoint parte lineas que
#: en el PDF caben. Va en absoluto y no en porcentaje: proporcional, las lineas
#: largas se desbordaban sobre la columna vecina.
HOLGURA_PT = 3.0
#: Cuando dos fragmentos consecutivos de una linea estan separados por mas de
#: esta fraccion del cuerpo menor, entre ellos habia un espacio. En las lineas
#: justificadas LaTeX no escribe glifo de espacio y pymupdf devuelve un
#: fragmento por palabra, de modo que sin esto salian 39 renglones con las
#: palabras pegadas: «Aesosesumanlasareasprotegidas». Medido sobre los 542
#: renglones de los dos mazos, cualquier valor entre 0,15 y 0,30 da el mismo
#: resultado y solo a partir de 0,40 se empiezan a perder espacios.
HUECO_ESPACIO = 0.20
#: Grosor minimo de un filete, en puntos del PDF. Los filetes de la marca son
#: trazos sin relleno cuyo rectangulo mide cero de alto.
GROSOR_MINIMO = 0.4


def _rgb(v):
    return RGBColor((v >> 16) & 255, (v >> 8) & 255, v & 255)


def _color(c, por_defecto=(0.5, 0.5, 0.5)):
    c = c or por_defecto
    return RGBColor(*[max(0, min(255, int(round(255 * v)))) for v in c[:3]])


def _normalizar(r):
    """Ordena el rectangulo y le da grosor minimo si mide cero por un lado.

    Un filete de 0,6 pt de grosor llega del PDF como un trazo sin relleno cuyo
    rectangulo tiene alto cero. El descarte anterior era `width <= 0.2 or
    height <= 0.2`, de modo que se perdian 23 filetes de verdad, entre ellos las
    tres reglas de columna de la lamina 3 comercial y la linea que separa la
    cabecera de la tabla de fuentes. Cambiar el `or` por un `and` no vale: en
    pymupdf `Rect.intersects()` devuelve False cuando uno de los dos rectangulos
    esta vacio, asi que la guarda que evita repintar lo que ya va dentro de una
    figura nunca atraparia un trazo de grosor cero, y volverian 113 formas
    sueltas encima de figuras que ya las contienen.
    """
    x0, x1 = sorted((r.x0, r.x1))
    y0, y1 = sorted((r.y0, r.y1))
    if x1 - x0 < GROSOR_MINIMO:
        x1 = x0 + GROSOR_MINIMO
    if y1 - y0 < GROSOR_MINIMO:
        y1 = y0 + GROSOR_MINIMO
    return pymupdf.Rect(x0, y0, x1, y1)


def _fusionar(cajas, margen=7.0):
    """Agrupa rectangulos que se tocan o casi, para separar figura de figura."""
    grupos = []
    for c in cajas:
        r = pymupdf.Rect(c)
        unido = None
        for g in grupos:
            if pymupdf.Rect(g[0] - margen, g[1] - margen,
                            g[2] + margen, g[3] + margen).intersects(r):
                unido = g
                break
        if unido is None:
            grupos.append([r.x0, r.y0, r.x1, r.y1])
        else:
            unido[0], unido[1] = min(unido[0], r.x0), min(unido[1], r.y0)
            unido[2], unido[3] = max(unido[2], r.x1), max(unido[3], r.y1)
    cambio = True
    while cambio:
        cambio = False
        for i in range(len(grupos)):
            for j in range(i + 1, len(grupos)):
                a, b = grupos[i], grupos[j]
                if pymupdf.Rect(*a).intersects(pymupdf.Rect(
                        b[0] - margen, b[1] - margen, b[2] + margen, b[3] + margen)):
                    a[0], a[1] = min(a[0], b[0]), min(a[1], b[1])
                    a[2], a[3] = max(a[2], b[2]), max(a[3], b[3])
                    grupos.pop(j)
                    cambio = True
                    break
            if cambio:
                break
    return [pymupdf.Rect(*g) for g in grupos]


def _lineas_latex(dic):
    """Las lineas de texto compuestas por LaTeX, con sus fragmentos."""
    for b in dic["blocks"]:
        if b["type"] != 0:
            continue
        for li in b["lines"]:
            spans = [sp for sp in li["spans"]
                     if not sp["font"].startswith("Arial") and sp["text"].strip()]
            if spans:
                yield li, spans


def _pagina_sin_texto(doc, n, rects):
    """Copia de la pagina con esos rectangulos de texto borrados.

    En la portada, el PNG del mapa trae grabadas en su borde las letras del
    final del titulo, y encima el cuadro de texto vuelve a escribirlo: en Arial
    no cae donde caia en NimbusSanL y se ve doble. Lo mismo pasa con los pies de
    las dos imagenes del catastro. No sirve recortar el rectangulo de la figura,
    porque el titulo entra por la izquierda y saldria un rectangulo invertido, y
    tampoco pintar de blanco encima, que da por hecho que debajo la figura es
    blanca. Se marca como redaccion cada linea y se aplica sin tocar imagenes ni
    trazos: desaparece la tinta del texto y la figura queda intacta.
    """
    copia = pymupdf.open()
    copia.insert_pdf(doc, from_page=n - 1, to_page=n - 1)
    pg = copia[0]
    for r in rects:
        pg.add_redact_annot(r, fill=False)
    pg.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                        text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    return copia, pg


def main():
    global PDF, SAL
    args = list(sys.argv[1:])
    # Si los dos primeros argumentos son rutas, son el PDF de entrada y el PPTX de
    # salida; lo que quede son numeros de lamina.
    while args and not args[0].lstrip("-").isdigit():
        ruta = pathlib.Path(args.pop(0))
        if ruta.suffix.lower() == ".pdf":
            PDF = ruta
        else:
            SAL = ruta
    quiere = {int(a) for a in args} or None
    TMP.mkdir(parents=True, exist_ok=True)
    for viejo in TMP.glob("*.png"):
        viejo.unlink()

    doc = pymupdf.open(PDF)
    pres = Presentation()
    pres.slide_width, pres.slide_height = Inches(PULGADAS[0]), Inches(PULGADAS[1])
    vacio = pres.slide_layouts[6]
    esc = pres.slide_width / doc[0].rect.width      # EMU por punto de PDF
    pt_por_pt = esc / 12700.0                       # de pt de PDF a pt de PPTX

    def emu(v):
        return Emu(int(v * esc))

    portada = None
    n_laminas = 0

    for n, pagina in enumerate(doc, 1):
        if quiere and n not in quiere:
            continue
        if portada is None:
            portada = pagina
        n_laminas += 1
        lam = pres.slides.add_slide(vacio)
        dic = pagina.get_text("dict")
        dibujos = pagina.get_drawings()

        # Que es figura y que es adorno de la lamina. Se agrupa todo lo dibujable
        # y despues se decide grupo a grupo: es figura si trae una imagen raster,
        # si trae texto de matplotlib, o si acumula muchos trazos, que es lo que
        # delata un grafico vectorial como el mapa de la portada. Lo demas son el
        # filete de marca, las lineas de tabla y el isotipo del pie.
        piezas = []
        for x in pagina.get_images(full=True):
            piezas.append((pymupdf.Rect(pagina.get_image_bbox(x)), "raster"))
        for b in dic["blocks"]:
            if b["type"] != 0:
                continue
            for li in b["lines"]:
                for sp in li["spans"]:
                    if sp["font"].startswith("Arial"):
                        piezas.append((pymupdf.Rect(sp["bbox"]), "arial"))
        # Lo que cruza la lamina de lado a lado no entra en el agrupamiento: son el
        # fondo blanco de la pagina y la banda del titulo, y al agrupar hacian de
        # puente entre elementos sin relacion, con lo que salia un solo grupo que
        # cubria la diapositiva entera.
        anchura, area = pagina.rect.width, pagina.rect.get_area()
        for dr in dibujos:
            r = _normalizar(pymupdf.Rect(dr["rect"]))
            if r.width < 0.92 * anchura:
                piezas.append((r, "trazo"))

        grupos = _fusionar([r for r, _ in piezas], margen=6.0)
        clases = []
        for g in grupos:
            dentro = [t for r, t in piezas if g.intersects(r)]
            es_fig = ("raster" in dentro or "arial" in dentro
                      or dentro.count("trazo") >= 8)
            clases.append(es_fig)
        figuras = [g for g, e in zip(grupos, clases) if e]

        # Las lineas de LaTeX que cruzan una figura se borran de la copia de la
        # que se rasteriza, para que no salgan cocidas dentro de la imagen y
        # repetidas encima como caja viva.
        pisadas = [pymupdf.Rect(li["bbox"]) for li, _ in _lineas_latex(dic)
                   if any(f.intersects(pymupdf.Rect(li["bbox"])) for f in figuras)]
        copia, fuente = (None, pagina)
        if pisadas and figuras:
            copia, fuente = _pagina_sin_texto(doc, n, pisadas)

        for k, f in enumerate(figuras, 1):
            f = pymupdf.Rect(f.x0 - 2, f.y0 - 2, f.x1 + 2, f.y1 + 2)
            ruta = TMP / f"{n:02d}_{k}.png"
            fuente.get_pixmap(clip=f, dpi=PPP_FIGURA).save(ruta)
            lam.shapes.add_picture(str(ruta), emu(f.x0), emu(f.y0),
                                   width=emu(f.width), height=emu(f.height))
        if copia is not None:
            copia.close()

        for dr in dibujos:
            r = _normalizar(pymupdf.Rect(dr["rect"]))
            # Lo que cae dentro de una figura ya esta en su imagen: repetirlo como
            # forma lo dibujaria dos veces.
            if any(f.intersects(r) for f in figuras):
                continue
            # El fondo de pagina y la banda del titulo no se dibujan: PowerPoint ya
            # da la diapositiva en blanco y repetirlos taparia el resto.
            if r.get_area() > 0.30 * area:
                continue
            fo = lam.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(r.x0), emu(r.y0),
                                      emu(r.width), emu(r.height))
            # Relleno y contorno son dos cosas distintas. Tomando `fill or color`
            # y rellenando siempre, las seis tarjetas de contorno de las laminas
            # de entregables salian pintadas de gris macizo con el color de su
            # propia linea y sin borde.
            relleno, trazo = dr.get("fill"), dr.get("color")
            if relleno:
                fo.fill.solid()
                fo.fill.fore_color.rgb = _color(relleno)
            else:
                fo.fill.background()
            if trazo:
                fo.line.color.rgb = _color(trazo)
                fo.line.width = Pt(round(max(dr.get("width") or 0.6, 0.3)
                                         * pt_por_pt, 2))
            else:
                fo.line.fill.background()
            fo.shadow.inherit = False

        # El texto de la lamina va SIEMPRE como cuadro editable. Una caja por
        # LINEA, no por bloque. Agrupando por bloque, las celdas de una tabla
        # caian en el mismo bloque y sus cajas se cruzaban entre si: en la
        # version tecnica salian 74 pares de texto solapados. La linea del PDF ya
        # viene con su rectangulo ajustado y dos lineas nunca se pisan.
        for li, spans in _lineas_latex(dic):
            r = pymupdf.Rect(li["bbox"])
            # Sin relleno vertical: la caja se cine al rectangulo de la linea.
            # Con cuatro puntos de mas, cada caja invadia la linea siguiente,
            # porque el paso entre lineas es menor que eso. PowerPoint no
            # recorta el texto que sobra, asi que ajustar no pierde nada.
            izq, arr = emu(r.x0 - 1.0), emu(r.y0 - 0.5)
            ancho = min(emu(r.width + HOLGURA_PT), pres.slide_width - izq)
            alto = min(emu(r.height + 1), pres.slide_height - arr)
            caja = lam.shapes.add_textbox(izq, arr, ancho, alto)
            marco = caja.text_frame
            marco.word_wrap = False
            marco.margin_left = marco.margin_right = 0
            marco.margin_top = marco.margin_bottom = 0
            par = marco.paragraphs[0]
            # La linea base del renglon es la del fragmento de mayor cuerpo: los
            # que van por encima son superindices, como el 2 de km2, y sin
            # marcarlos como tales caian a la altura del texto normal.
            base = max(spans, key=lambda s: s["size"])["origin"][1]
            previo = None
            for sp in spans:
                texto = sp["text"]
                if previo is not None:
                    hueco = sp["bbox"][0] - previo["bbox"][2]
                    if (hueco > HUECO_ESPACIO * min(sp["size"], previo["size"])
                            and not previo["text"].endswith(" ")
                            and not texto.startswith(" ")):
                        texto = " " + texto
                run = par.add_run()
                run.text = texto
                run.font.name = TIPO
                run.font.size = Pt(round(sp["size"] * pt_por_pt, 1))
                run.font.bold = "Bold" in sp["font"]
                run.font.color.rgb = _rgb(sp["color"])
                alza = base - sp["origin"][1]
                if sp["size"] > 0 and alza > 0.15 * sp["size"]:
                    run.font._rPr.set("baseline",
                                      str(int(round(alza / sp["size"] * 100000))))
                previo = sp

        print(f"  lamina {n}: {len(figuras)} figura(s), "
              f"{len(lam.shapes)} elementos")

    titulo = "Prospectos Solares, versión editable"
    pptx_comun.fijar_propiedades(pres, titulo, datetime.datetime.now())
    pres.save(SAL)
    pptx_comun.sanear(SAL, portada, n_laminas, titulo)
    print(f"\n{SAL}\n{SAL.stat().st_size / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
